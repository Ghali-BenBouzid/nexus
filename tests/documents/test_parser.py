import io

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.documents.parser import ParseError, parse

_TEXT = (
    "Heat pumps move heat instead of making it, which is why they can deliver "
    "more energy than they consume."
)
_CSS = "table {border-collapse: collapse} td, th {border: 1px solid #000; padding: 6px}"


def _pdf(html: str | None = None, pages: int = 1) -> bytes:
    """A real PDF, written with pymupdf. ``html`` is laid out on every page."""
    document = pymupdf.open()
    for index in range(pages):
        page = document.new_page()
        body = html or f"<p>{_TEXT} Page {index + 1}.</p>"
        page.insert_htmlbox(pymupdf.Rect(50, 50, 545, 750), body, css=_CSS)
    data = document.tobytes()
    document.close()
    return data


def _scanned_pdf(pages: int = 2) -> bytes:
    """Pages with a picture and no text: what a scan looks like to a parser."""
    document = pymupdf.open()
    pixmap = pymupdf.Pixmap(pymupdf.csGRAY, pymupdf.IRect(0, 0, 200, 200), False)
    pixmap.clear_with(128)
    for _ in range(pages):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(50, 50, 545, 750), pixmap=pixmap)
    data = document.tobytes()
    document.close()
    return data


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(added.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_pdf_gives_its_text_and_page_count() -> None:
    parsed = parse("report.pdf", _pdf(pages=2))

    assert "Heat pumps move heat" in parsed.text
    assert "Page 2." in parsed.text
    assert parsed.pages == 2


def test_a_pdf_keeps_its_headings_and_tables_as_markdown() -> None:
    # What the agents will work from: a section can be quoted, a row read back.
    html = """
    <h1>Heat pump performance</h1>
    <p>Efficiency falls in the cold, but modern units keep working.</p>
    <table><tr><th>Outdoor temperature</th><th>COP</th></tr>
    <tr><td>7 C</td><td>4.1</td></tr><tr><td>-15 C</td><td>2.2</td></tr></table>
    """

    text = parse("performance.pdf", _pdf(html)).text

    assert "# **Heat pump performance**" in text
    assert "|**Outdoor temperature**|**COP**|" in text
    assert "|7 C|4.1|" in text


def test_a_scanned_pdf_is_refused_and_says_why() -> None:
    # It parses fine and yields almost nothing: the failure a user must understand.
    with pytest.raises(ParseError, match="scanned"):
        parse("scan.pdf", _scanned_pdf())


def test_a_damaged_pdf_is_refused() -> None:
    with pytest.raises(ParseError, match="could not be read"):
        parse("broken.pdf", b"%PDF-1.4\nnot really a pdf")


def test_a_word_document_gives_its_paragraphs_and_tables() -> None:
    data = _docx(
        ["The rules of the tender are set out below." * 2],
        table=[["Criterion", "Weight"], ["Price", "40%"]],
    )

    parsed = parse("tender.docx", data)

    assert "The rules of the tender" in parsed.text
    assert "Criterion | Weight" in parsed.text
    assert "Price | 40%" in parsed.text
    assert parsed.pages is None


def test_plain_text_survives_a_stray_byte() -> None:
    parsed = parse("notes.md", b"# Notes\n\nSomething worth keeping." + b"\xff")

    assert parsed.text.startswith("# Notes")


def test_an_unsupported_type_is_named_in_the_error() -> None:
    with pytest.raises(ParseError, match="not a file type"):
        parse("slides.pptx", b"whatever")


def test_a_file_with_no_readable_text_is_refused() -> None:
    with pytest.raises(ParseError, match="No text"):
        parse("empty.txt", b"  \n ")
